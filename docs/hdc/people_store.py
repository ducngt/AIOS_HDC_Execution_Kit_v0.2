"""Foundation Identity & People store for the AIOS integrated prototype.

Prototype persistence is SQLite. Uploaded institutional workbooks are parsed at
runtime and are not committed to the source repository.
"""
from __future__ import annotations

import base64
import hashlib
import secrets
import io
import json
import os
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB = ROOT / "runtime" / "aios-prototype.sqlite3"


def db_path() -> Path:
    return Path(os.getenv("AIOS_DB_PATH", str(DEFAULT_DB)))


def connect() -> sqlite3.Connection:
    path = db_path().expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(path, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA journal_mode = WAL")
    con.execute("PRAGMA synchronous = NORMAL")
    con.execute("PRAGMA busy_timeout = 30000")
    ensure_schema(con)
    return con


def ensure_schema(con: sqlite3.Connection) -> None:
    con.executescript(
        """
        PRAGMA foreign_keys = ON;
        CREATE TABLE IF NOT EXISTS organizations(
          org_code TEXT PRIMARY KEY,
          parent_code TEXT,
          org_type TEXT,
          name TEXT NOT NULL,
          short_name TEXT,
          active INTEGER NOT NULL DEFAULT 1,
          raw_json TEXT
        );
        CREATE TABLE IF NOT EXISTS positions(
          id TEXT PRIMARY KEY,
          org_code TEXT,
          reference_code TEXT,
          level TEXT,
          title TEXT NOT NULL,
          headcount REAL,
          position_type TEXT,
          raw_json TEXT
        );
        CREATE TABLE IF NOT EXISTS people(
          person_id TEXT PRIMARY KEY,
          source_key TEXT UNIQUE,
          person_kind TEXT NOT NULL,
          full_name TEXT NOT NULL,
          family_middle TEXT,
          given_name TEXT,
          gender TEXT,
          birth_date TEXT,
          national_id TEXT,
          email TEXT,
          phone TEXT,
          status TEXT,
          primary_org_code TEXT,
          student_major_cohort_code TEXT,
          class_code TEXT,
          raw_json TEXT,
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS person_org_assignments(
          id TEXT PRIMARY KEY,
          person_id TEXT NOT NULL,
          org_code TEXT NOT NULL,
          assignment_type TEXT NOT NULL,
          active INTEGER NOT NULL DEFAULT 1,
          UNIQUE(person_id, org_code, assignment_type),
          FOREIGN KEY(person_id) REFERENCES people(person_id)
        );
        CREATE TABLE IF NOT EXISTS student_cohorts(
          cohort_code TEXT PRIMARY KEY,
          name TEXT,
          start_year TEXT,
          education_level_code TEXT,
          training_mode_code TEXT,
          raw_json TEXT
        );
        CREATE TABLE IF NOT EXISTS major_cohorts(
          major_cohort_code TEXT PRIMARY KEY,
          cohort_code TEXT,
          major_code TEXT,
          major_name TEXT,
          name TEXT,
          curriculum_code TEXT,
          campus_code TEXT,
          curriculum_type_code TEXT,
          raw_json TEXT
        );
        CREATE TABLE IF NOT EXISTS admin_classes(
          class_code TEXT PRIMARY KEY,
          major_cohort_code TEXT,
          max_size INTEGER,
          raw_json TEXT
        );
        CREATE TABLE IF NOT EXISTS accounts(
          account_id TEXT PRIMARY KEY,
          person_id TEXT UNIQUE,
          username TEXT UNIQUE NOT NULL,
          display_name TEXT NOT NULL,
          email TEXT,
          status TEXT NOT NULL DEFAULT 'active',
          locale TEXT NOT NULL DEFAULT 'vi',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          FOREIGN KEY(person_id) REFERENCES people(person_id)
        );
        CREATE TABLE IF NOT EXISTS role_assignments(
          id TEXT PRIMARY KEY,
          account_id TEXT NOT NULL,
          role_code TEXT NOT NULL,
          scope TEXT,
          active INTEGER NOT NULL DEFAULT 1,
          UNIQUE(account_id, role_code, scope),
          FOREIGN KEY(account_id) REFERENCES accounts(account_id)
        );
        CREATE TABLE IF NOT EXISTS imports(
          import_id TEXT PRIMARY KEY,
          filename TEXT NOT NULL,
          detected_type TEXT NOT NULL,
          status TEXT NOT NULL,
          row_count INTEGER NOT NULL,
          created_count INTEGER NOT NULL,
          updated_count INTEGER NOT NULL,
          warning_count INTEGER NOT NULL,
          error_count INTEGER NOT NULL,
          report_json TEXT NOT NULL,
          created_at TEXT NOT NULL
        );
        """
    )
    # Lightweight prototype migrations for account authentication.
    cols = {r[1] for r in con.execute("PRAGMA table_info(accounts)").fetchall()}
    for name, ddl in {
        "password_hash": "ALTER TABLE accounts ADD COLUMN password_hash TEXT",
        "password_salt": "ALTER TABLE accounts ADD COLUMN password_salt TEXT",
        "must_change_password": "ALTER TABLE accounts ADD COLUMN must_change_password INTEGER NOT NULL DEFAULT 1",
    }.items():
        if name not in cols:
            con.execute(ddl)
    con.commit()


def now() -> str:
    return datetime.utcnow().replace(microsecond=0).isoformat() + "Z"


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        s = value.strip()
        return s or None
    return value


def row_dict(headers: list[Any], row: Iterable[Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for h, v in zip(headers, row):
        if h is None:
            continue
        out[str(h).strip()] = clean(v)
    return out


def parse_date(*values: Any) -> str | None:
    for value in values:
        value = clean(value)
        if not value:
            continue
        if isinstance(value, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(value, fmt).date().isoformat()
                except ValueError:
                    pass
            if len(value) >= 10 and value[4:5] == "-":
                return value[:10]
    return None


def workbook_from_base64(data_b64: str):
    raw = base64.b64decode(data_b64, validate=True)
    return load_workbook(io.BytesIO(raw), read_only=True, data_only=False)


def detect_workbook(wb) -> str:
    names = set(wb.sheetnames)
    sample = []
    for ws in wb.worksheets:
        if not ws.title.startswith("Data"):
            continue
        sample = [clean(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        break
    h = {str(x) for x in sample if x is not None}
    if "Mã cán bộ(*)" in h:
        return "personnel"
    if "Mã đơn vị" in h and "Tên vị trí chức danh" in h:
        return "positions"
    if "Loại phòng ban" in h and "Mã đơn vị" in h:
        return "organizations"
    if "Mã khóa ngành" in h and "Mã ngành đào tạo" in h:
        return "major_cohorts"
    if "Mã khóa sinh viên" in h:
        return "student_cohorts"
    if "Tên lớp" in h and "Mã khóa ngành" in h:
        return "admin_classes"
    if "Mã người học" in h:
        return "students"
    return "unknown"


@dataclass
class ImportReport:
    import_id: str
    filename: str
    detected_type: str
    row_count: int = 0
    created_count: int = 0
    updated_count: int = 0
    warnings: list[str] | None = None
    errors: list[str] | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "import_id": self.import_id,
            "filename": self.filename,
            "detected_type": self.detected_type,
            "row_count": self.row_count,
            "created_count": self.created_count,
            "updated_count": self.updated_count,
            "warning_count": len(self.warnings or []),
            "error_count": len(self.errors or []),
            "warnings": self.warnings or [],
            "errors": self.errors or [],
        }


def upsert_person(con, *, source_key: str, person_kind: str, full_name: str,
                  family_middle: str | None = None, given_name: str | None = None,
                  gender: str | None = None, birth_date: str | None = None,
                  national_id: str | None = None, email: str | None = None,
                  phone: str | None = None, status: str | None = None,
                  primary_org_code: str | None = None,
                  student_major_cohort_code: str | None = None,
                  class_code: str | None = None, raw: dict | None = None) -> tuple[str, bool]:
    row = con.execute("SELECT person_id FROM people WHERE source_key=?", (source_key,)).fetchone()
    ts = now()
    if row:
        pid = row["person_id"]
        con.execute(
            """UPDATE people SET person_kind=?,full_name=?,family_middle=?,given_name=?,gender=?,birth_date=?,national_id=?,email=?,phone=?,status=?,primary_org_code=?,student_major_cohort_code=?,class_code=?,raw_json=?,updated_at=? WHERE person_id=?""",
            (person_kind, full_name, family_middle, given_name, gender, birth_date, national_id, email, phone, status, primary_org_code, student_major_cohort_code, class_code, json.dumps(raw or {}, ensure_ascii=False, default=str), ts, pid),
        )
        return pid, False
    pid = "PER-" + uuid.uuid4().hex[:16].upper()
    con.execute(
        """INSERT INTO people(person_id,source_key,person_kind,full_name,family_middle,given_name,gender,birth_date,national_id,email,phone,status,primary_org_code,student_major_cohort_code,class_code,raw_json,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (pid, source_key, person_kind, full_name, family_middle, given_name, gender, birth_date, national_id, email, phone, status, primary_org_code, student_major_cohort_code, class_code, json.dumps(raw or {}, ensure_ascii=False, default=str), ts, ts),
    )
    return pid, True


def import_workbook(filename: str, data_b64: str) -> dict[str, Any]:
    wb = workbook_from_base64(data_b64)
    kind = detect_workbook(wb)
    report = ImportReport("IMP-" + uuid.uuid4().hex[:12].upper(), filename, kind, warnings=[], errors=[])
    if kind == "unknown":
        report.errors.append("Không nhận diện được cấu trúc workbook.")
        return report.as_dict()
    with connect() as con:
        try:
            dispatch = {
                "organizations": import_organizations,
                "positions": import_positions,
                "personnel": import_personnel,
                "student_cohorts": import_student_cohorts,
                "major_cohorts": import_major_cohorts,
                "admin_classes": import_admin_classes,
                "students": import_students,
            }
            dispatch[kind](con, wb, report)
            con.commit()
        except Exception as exc:
            con.rollback()
            report.errors.append(f"Import thất bại: {exc}")
        d = report.as_dict()
        con.execute(
            """INSERT OR REPLACE INTO imports(import_id,filename,detected_type,status,row_count,created_count,updated_count,warning_count,error_count,report_json,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (report.import_id, filename, kind, "failed" if d["error_count"] else "completed", d["row_count"], d["created_count"], d["updated_count"], d["warning_count"], d["error_count"], json.dumps(d, ensure_ascii=False), now()),
        )
        con.commit()
    return report.as_dict()


def data_sheet(wb):
    for ws in wb.worksheets:
        if ws.title.startswith("Data"):
            return ws
    return wb.worksheets[0]


def import_organizations(con, wb, report):
    ws = data_sheet(wb); headers = [clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        d=row_dict(headers,row); code=clean(d.get("Mã đơn vị")); name=clean(d.get("Tên"))
        if not code or not name: continue
        report.row_count += 1
        existed=con.execute("SELECT 1 FROM organizations WHERE org_code=?",(str(code),)).fetchone() is not None
        con.execute("""INSERT INTO organizations(org_code,parent_code,org_type,name,short_name,active,raw_json) VALUES(?,?,?,?,?,?,?) ON CONFLICT(org_code) DO UPDATE SET parent_code=excluded.parent_code,org_type=excluded.org_type,name=excluded.name,short_name=excluded.short_name,active=excluded.active,raw_json=excluded.raw_json""",
                    (str(code), clean(d.get("Mã đơn vị cấp trên")), clean(d.get("Loại phòng ban")), str(name), clean(d.get("Tên viết tắt")), 1 if d.get("Hiển thị") in (None,1,"1",True) else 0, json.dumps(d,ensure_ascii=False,default=str)))
        report.updated_count += int(existed); report.created_count += int(not existed)


def import_positions(con, wb, report):
    ws=data_sheet(wb); headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]; current_org=None
    for idx,row in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
        d=row_dict(headers,row); current_org=clean(d.get("Mã đơn vị")) or current_org; title=clean(d.get("Tên vị trí chức danh"))
        if not title: continue
        report.row_count += 1
        ref=str(clean(d.get("Mã vị trí chức danh tham khảo")) or "")
        pid=f"POS-{current_org or 'NA'}-{ref}-{idx}"
        existed=con.execute("SELECT 1 FROM positions WHERE id=?",(pid,)).fetchone() is not None
        con.execute("""INSERT INTO positions(id,org_code,reference_code,level,title,headcount,position_type,raw_json) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET org_code=excluded.org_code,reference_code=excluded.reference_code,level=excluded.level,title=excluded.title,headcount=excluded.headcount,position_type=excluded.position_type,raw_json=excluded.raw_json""",
                    (pid,current_org,ref,clean(d.get("Cấp chức vụ")),str(title),clean(d.get("Số lượng người làm việc")),clean(d.get("Loại")),json.dumps(d,ensure_ascii=False,default=str)))
        report.updated_count += int(existed); report.created_count += int(not existed)


def import_personnel(con, wb, report):
    ws=data_sheet(wb); headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]; current_pid=None
    for row in ws.iter_rows(min_row=2, values_only=True):
        d=row_dict(headers,row); staff=clean(d.get("Mã cán bộ(*)")); org=clean(d.get("Mã đơn vị(*)"))
        if not staff:
            if current_pid and org:
                con.execute("INSERT OR IGNORE INTO person_org_assignments(id,person_id,org_code,assignment_type) VALUES(?,?,?,?)",("POA-"+uuid.uuid4().hex[:16].upper(),current_pid,str(org),"concurrent"))
            continue
        family=clean(d.get("Họ đệm(*)")) or ""; given=clean(d.get("Tên(*)")) or ""; full=(str(family)+" "+str(given)).strip()
        if not full: continue
        report.row_count += 1
        pid,created=upsert_person(con,source_key="staff:"+str(staff),person_kind="staff",full_name=full,family_middle=str(family) or None,given_name=str(given) or None,gender=clean(d.get("Giới tính(*)")),birth_date=parse_date(d.get("Ngày sinh")),national_id=str(clean(d.get("CCCD/CMND(*)")) or "") or None,email=clean(d.get("Email cán bộ")) or clean(d.get("Email")),phone=str(clean(d.get("Điện thoại")) or "") or None,status=clean(d.get("Trạng thái làm việc")),primary_org_code=str(org) if org else None,raw=d)
        current_pid=pid
        if org: con.execute("INSERT OR IGNORE INTO person_org_assignments(id,person_id,org_code,assignment_type) VALUES(?,?,?,?)",("POA-"+uuid.uuid4().hex[:16].upper(),pid,str(org),"primary"))
        report.created_count += int(created); report.updated_count += int(not created)


def import_student_cohorts(con, wb, report):
    ws=data_sheet(wb); headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        d=row_dict(headers,row); code=clean(d.get("Mã khóa sinh viên"))
        if not code: continue
        report.row_count+=1; existed=con.execute("SELECT 1 FROM student_cohorts WHERE cohort_code=?",(str(code),)).fetchone() is not None
        con.execute("""INSERT INTO student_cohorts(cohort_code,name,start_year,education_level_code,training_mode_code,raw_json) VALUES(?,?,?,?,?,?) ON CONFLICT(cohort_code) DO UPDATE SET name=excluded.name,start_year=excluded.start_year,education_level_code=excluded.education_level_code,training_mode_code=excluded.training_mode_code,raw_json=excluded.raw_json""",(str(code),clean(d.get("Tên khóa sinh viên")),str(clean(d.get("Năm học bắt đầu")) or "") or None,str(clean(d.get("Mã trình độ đào tạo")) or "") or None,str(clean(d.get("Mã hình thức đào tạo")) or "") or None,json.dumps(d,ensure_ascii=False,default=str)))
        report.created_count+=int(not existed); report.updated_count+=int(existed)


def import_major_cohorts(con, wb, report):
    ws=data_sheet(wb); headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        d=row_dict(headers,row); code=clean(d.get("Mã khóa ngành"))
        if not code: continue
        report.row_count+=1; existed=con.execute("SELECT 1 FROM major_cohorts WHERE major_cohort_code=?",(str(code),)).fetchone() is not None
        con.execute("""INSERT INTO major_cohorts(major_cohort_code,cohort_code,major_code,major_name,name,curriculum_code,campus_code,curriculum_type_code,raw_json) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(major_cohort_code) DO UPDATE SET cohort_code=excluded.cohort_code,major_code=excluded.major_code,major_name=excluded.major_name,name=excluded.name,curriculum_code=excluded.curriculum_code,campus_code=excluded.campus_code,curriculum_type_code=excluded.curriculum_type_code,raw_json=excluded.raw_json""",(str(code),str(clean(d.get("Mã khóa sinh viên")) or "") or None,str(clean(d.get("Mã ngành đào tạo")) or "") or None,clean(d.get("Ngành đào tạo")),clean(d.get("Tên")),str(clean(d.get("Mã chương trình đào tạo")) or "") or None,str(clean(d.get("Mã cơ sở ĐT")) or "") or None,str(clean(d.get("Mã tính chất CTĐT")) or "") or None,json.dumps(d,ensure_ascii=False,default=str)))
        report.created_count+=int(not existed); report.updated_count+=int(existed)


def import_admin_classes(con, wb, report):
    ws=data_sheet(wb); headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    for row in ws.iter_rows(min_row=2, values_only=True):
        d=row_dict(headers,row); code=clean(d.get("Tên lớp"))
        if not code: continue
        report.row_count+=1; existed=con.execute("SELECT 1 FROM admin_classes WHERE class_code=?",(str(code),)).fetchone() is not None
        con.execute("""INSERT INTO admin_classes(class_code,major_cohort_code,max_size,raw_json) VALUES(?,?,?,?) ON CONFLICT(class_code) DO UPDATE SET major_cohort_code=excluded.major_cohort_code,max_size=excluded.max_size,raw_json=excluded.raw_json""",(str(code),str(clean(d.get("Mã khóa ngành")) or "") or None,clean(d.get("Sĩ số tối đa")),json.dumps(d,ensure_ascii=False,default=str)))
        report.created_count+=int(not existed); report.updated_count+=int(existed)


def import_students(con, wb, report):
    for ws in wb.worksheets:
        if not ws.title.startswith("Data"):
            continue
        headers=[clean(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
        if "Mã người học" not in headers: continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            d=row_dict(headers,row); sid=clean(d.get("Mã người học")); full=clean(d.get("Họ và tên"))
            if not sid or not full: continue
            report.row_count+=1
            pid,created=upsert_person(con,source_key="student:"+str(sid),person_kind="student",full_name=str(full),family_middle=clean(d.get("Họ đệm")),given_name=clean(d.get("Tên")),gender=clean(d.get("Giới tính")),birth_date=parse_date(d.get("Ngày sinh (YYYY-MM-DD)"),d.get("Ngày sinh (DD/MM/YYYY)")),national_id=str(clean(d.get("CMT/CCCD")) or "") or None,email=clean(d.get("Email")),phone=str(clean(d.get("Số điện thoại")) or "") or None,status=clean(d.get("Trạng thái học")) or ws.title.replace("Data__", ""),student_major_cohort_code=str(clean(d.get("Mã khóa ngành")) or "") or None,class_code=str(clean(d.get("Mã lớp hành chính")) or "") or None,raw=d)
            report.created_count+=int(created); report.updated_count+=int(not created)


def stats() -> dict[str, int]:
    with connect() as con:
        def c(sql): return int(con.execute(sql).fetchone()[0])
        return {
            "people": c("SELECT COUNT(*) FROM people"),
            "staff": c("SELECT COUNT(*) FROM people WHERE person_kind='staff'"),
            "students": c("SELECT COUNT(*) FROM people WHERE person_kind='student'"),
            "organizations": c("SELECT COUNT(*) FROM organizations"),
            "positions": c("SELECT COUNT(*) FROM positions"),
            "student_cohorts": c("SELECT COUNT(*) FROM student_cohorts"),
            "major_cohorts": c("SELECT COUNT(*) FROM major_cohorts"),
            "classes": c("SELECT COUNT(*) FROM admin_classes"),
            "accounts": c("SELECT COUNT(*) FROM accounts"),
        }


def list_people(kind: str | None = None, query: str | None = None, limit: int = 100) -> list[dict[str, Any]]:
    sql="SELECT person_id,source_key,person_kind,full_name,email,phone,status,primary_org_code,student_major_cohort_code,class_code FROM people WHERE 1=1"; args=[]
    if kind: sql += " AND person_kind=?"; args.append(kind)
    if query:
        sql += " AND (full_name LIKE ? OR source_key LIKE ? OR email LIKE ?)"; q=f"%{query}%"; args.extend([q,q,q])
    sql += " ORDER BY full_name LIMIT ?"; args.append(max(1,min(limit,500)))
    with connect() as con: return [dict(r) for r in con.execute(sql,args).fetchall()]


def list_accounts(limit: int = 200) -> list[dict[str, Any]]:
    with connect() as con:
        rows=con.execute("""SELECT a.account_id,a.person_id,a.username,a.display_name,a.email,a.status,a.locale,p.person_kind,p.source_key FROM accounts a LEFT JOIN people p ON p.person_id=a.person_id ORDER BY a.display_name LIMIT ?""",(max(1,min(limit,500)),)).fetchall()
        out=[]
        for r in rows:
            d=dict(r); d["roles"]=[dict(x) for x in con.execute("SELECT role_code,scope,active FROM role_assignments WHERE account_id=? ORDER BY role_code",(r["account_id"],)).fetchall()]; out.append(d)
        return out


def _password_digest(password: str, salt_hex: str) -> str:
    return hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), 180000).hex()


def set_password(con: sqlite3.Connection, account_id: str, password: str) -> None:
    if len(password) < 6:
        raise ValueError("Mật khẩu tạm thời phải có ít nhất 6 ký tự")
    salt = secrets.token_hex(16)
    digest = _password_digest(password, salt)
    con.execute("UPDATE accounts SET password_hash=?,password_salt=?,must_change_password=1,updated_at=? WHERE account_id=?", (digest, salt, now(), account_id))


def authenticate(username: str, password: str) -> dict[str, Any] | None:
    with connect() as con:
        row = con.execute("SELECT * FROM accounts WHERE lower(username)=lower(?) OR lower(email)=lower(?)", (username, username)).fetchone()
        if not row or row["status"] != "active" or not row["password_hash"] or not row["password_salt"]:
            return None
        if not secrets.compare_digest(row["password_hash"], _password_digest(password, row["password_salt"])):
            return None
        roles=[dict(x) for x in con.execute("SELECT role_code,scope,active FROM role_assignments WHERE account_id=? AND active=1 ORDER BY role_code",(row["account_id"],)).fetchall()]
        return {"account_id":row["account_id"],"person_id":row["person_id"],"name":row["display_name"],"email":row["email"],"username":row["username"],"status":row["status"],"locale":row["locale"],"must_change_password":bool(row["must_change_password"]),"roles":roles}


def save_account(payload: dict[str, Any]) -> dict[str, Any]:
    person_id=payload.get("person_id"); username=str(payload.get("username") or "").strip(); display=str(payload.get("display_name") or "").strip()
    if not username or not display: raise ValueError("username và display_name là bắt buộc")
    ts=now(); account_id=str(payload.get("account_id") or ("ACC-"+uuid.uuid4().hex[:14].upper()))
    with connect() as con:
        existing=con.execute("SELECT 1 FROM accounts WHERE account_id=?",(account_id,)).fetchone()
        if existing:
            con.execute("UPDATE accounts SET person_id=?,username=?,display_name=?,email=?,status=?,locale=?,updated_at=? WHERE account_id=?",(person_id,username,display,payload.get("email"),payload.get("status","active"),payload.get("locale","vi"),ts,account_id))
        else:
            con.execute("INSERT INTO accounts(account_id,person_id,username,display_name,email,status,locale,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",(account_id,person_id,username,display,payload.get("email"),payload.get("status","active"),payload.get("locale","vi"),ts,ts))
        if "roles" in payload:
            con.execute("DELETE FROM role_assignments WHERE account_id=?",(account_id,))
            for role in payload.get("roles") or []:
                code=str(role.get("role_code") or "").strip()
                if not code: continue
                con.execute("INSERT INTO role_assignments(id,account_id,role_code,scope,active) VALUES(?,?,?,?,?)",("RA-"+uuid.uuid4().hex[:14].upper(),account_id,code,role.get("scope"),1 if role.get("active",True) else 0))
        temporary_password = payload.get("temporary_password")
        if temporary_password:
            set_password(con, account_id, str(temporary_password))
        con.commit()
    return {"account_id":account_id,"saved":True,"password_updated":bool(payload.get("temporary_password"))}


def toggle_account(account_id: str, status: str) -> dict[str, Any]:
    if status not in {"active","locked","disabled"}: raise ValueError("invalid status")
    with connect() as con:
        cur=con.execute("UPDATE accounts SET status=?,updated_at=? WHERE account_id=?",(status,now(),account_id)); con.commit()
        if cur.rowcount != 1: raise ValueError("account not found")
    return {"account_id":account_id,"status":status}


def list_imports(limit: int = 100) -> list[dict[str, Any]]:
    with connect() as con:
        rows = con.execute(
            "SELECT import_id,filename,detected_type,status,row_count,created_count,updated_count,warning_count,error_count,created_at FROM imports ORDER BY created_at DESC LIMIT ?",
            (max(1, min(limit, 500)),),
        ).fetchall()
        return [dict(r) for r in rows]


def database_info() -> dict[str, Any]:
    path = db_path().expanduser().resolve()
    return {
        "path": str(path),
        "exists": path.exists(),
        "size_bytes": path.stat().st_size if path.exists() else 0,
        "persistence": "sqlite",
    }


def save_data_record(payload: dict[str, Any]) -> dict[str, Any]:
    record_id = str(payload.get("record_id") or ("REC-" + uuid.uuid4().hex[:16].upper()))
    record_type = str(payload.get("type") or "Document")
    zone = str(payload.get("zone") or "D3")
    value = str(payload.get("value") or "").strip()
    if not value:
        raise ValueError("record value is required")
    ts = now()
    with connect() as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS data_records(
              record_id TEXT PRIMARY KEY, record_type TEXT NOT NULL, zone TEXT NOT NULL, value TEXT NOT NULL,
              validation_state TEXT NOT NULL DEFAULT 'candidate', created_by TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            )"""
        )
        con.execute(
            """INSERT INTO data_records(record_id,record_type,zone,value,validation_state,created_by,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(record_id) DO UPDATE SET record_type=excluded.record_type,zone=excluded.zone,value=excluded.value,validation_state=excluded.validation_state,created_by=excluded.created_by,updated_at=excluded.updated_at""",
            (record_id,record_type,zone,value,str(payload.get("validation") or "candidate"),payload.get("created_by"),ts,ts),
        )
        con.commit()
    return {"record_id":record_id,"saved":True}


def list_data_records(limit: int = 100) -> list[dict[str, Any]]:
    with connect() as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS data_records(
              record_id TEXT PRIMARY KEY, record_type TEXT NOT NULL, zone TEXT NOT NULL, value TEXT NOT NULL,
              validation_state TEXT NOT NULL DEFAULT 'candidate', created_by TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            )"""
        )
        return [dict(r) for r in con.execute(
            "SELECT * FROM data_records ORDER BY updated_at DESC LIMIT ?", (max(1,min(limit,500)),)
        ).fetchall()]
